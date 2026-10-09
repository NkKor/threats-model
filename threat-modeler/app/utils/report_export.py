"""Формирование итогового отчёта в форматах XLSX и DOCX.

Состав таблиц, заголовки, порядок и ширину колонок задаёт файл
``data/report-rules.txt`` (см. ``app/core/report_rules.py``), поэтому правка
правил в блокноте сразу меняет и предпросмотр, и выгрузку.
"""

from __future__ import annotations

from io import BytesIO
from typing import Any, Dict, List, Sequence

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.core.report_rules import Column, ReportRules

CELL_FONT = "Times New Roman"


# --------------------------------------------------------------------------- #
# Вспомогательные данные
# --------------------------------------------------------------------------- #

def profile_lines(report, rules: ReportRules) -> List[str]:
    """Строки с параметрами опроса (для DOCX)."""
    from app.core.reference_loader import ReferenceLoader

    types = ReferenceLoader.get_system_types()
    selected_types = [
        types.get(code, {}).get("name", code) for code in report.profile.system.system_types
    ]
    return [
        "Типы информационных систем: " + (", ".join(selected_types) or "не указаны"),
        "Объекты воздействия: "
        + (", ".join(report.statistics.get("selected_objects_names") or []) or "не указаны"),
        "Виды воздействия: " + (", ".join(report.profile.selected_impacts) or "не ограничено"),
        f"Всего УБИ в перечне: {report.total_threats} из {report.total_corpus}",
    ]


def _header_rows(table: Dict[str, Any]):
    """Строки шапки таблицы (с учётом групповых заголовков)."""
    return table.get("header_rows") or table["table"].header_rows()


# --------------------------------------------------------------------------- #
# XLSX
# --------------------------------------------------------------------------- #

def _write_sheet(worksheet, table: Dict[str, Any]) -> None:
    """Заполнить лист по правилам таблицы (в том числе объединённой шапкой)."""
    columns: List[Column] = table["table"].columns
    rows: List[List[str]] = table["rows"]
    header_rows = _header_rows(table)
    border = Side(style="thin")
    full_border = Border(left=border, right=border, top=border, bottom=border)
    header_font = Font(bold=True, size=11, name=CELL_FONT)
    cell_font = Font(size=10, name=CELL_FONT)
    header_fill = PatternFill(start_color="E0E0E0", end_color="E0E0E0", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    worksheet.cell(row=1, column=1, value=table["title"]).font = Font(bold=True, size=12, name=CELL_FONT)
    worksheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(len(columns), 1))

    # Шапка: одна или две строки, ячейки при необходимости объединяются
    for row_offset, header_row in enumerate(header_rows):
        for header_cell in header_row:
            col = header_cell.col + 1
            cell = worksheet.cell(row=2 + row_offset, column=col)
            cell.value = header_cell.text
            cell.font = header_font
            cell.fill = header_fill
            cell.border = full_border
            cell.alignment = header_align
            if header_cell.colspan > 1 or header_cell.rowspan > 1:
                worksheet.merge_cells(
                    start_row=2 + row_offset,
                    start_column=col,
                    end_row=2 + row_offset + header_cell.rowspan - 1,
                    end_column=col + header_cell.colspan - 1,
                )

    data_start = 2 + len(header_rows)
    for row_index, row in enumerate(rows, start=data_start):
        for column_index, value in enumerate(row, start=1):
            column = columns[column_index - 1]
            cell = worksheet.cell(row=row_index, column=column_index, value=value)
            cell.font = cell_font
            cell.border = full_border
            cell.alignment = Alignment(horizontal=column.align, vertical="top", wrap_text=True)

    for index, column in enumerate(columns, start=1):
        worksheet.column_dimensions[get_column_letter(index)].width = column.width
    worksheet.freeze_panes = f"A{data_start}"


def build_xlsx(report, tables: Dict[str, Dict[str, Any]], rules: ReportRules) -> bytes:
    """Собрать XLSX-отчёт по правилам."""
    workbook = Workbook()
    workbook.remove(workbook.active)
    for key, table in tables.items():
        worksheet = workbook.create_sheet(title=f"Таблица {key[-1]}")
        _write_sheet(worksheet, table)
    if not workbook.sheetnames:  # pragma: no cover - защита от пустых правил
        workbook.create_sheet(title="Отчёт")
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


# --------------------------------------------------------------------------- #
# DOCX
# --------------------------------------------------------------------------- #

def _style_table(table, header_height: int = 1) -> None:
    """Оформление таблицы Word: границы, шрифт, выравнивание."""
    table.style = "Table Grid"
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.name = CELL_FONT
                    run.font.size = Pt(9)
                    run.font.bold = row_index < header_height
            cell.vertical_alignment = 1


def _add_heading(document: Document, text: str, size: int = 13) -> None:
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    run.font.name = CELL_FONT
    run.font.size = Pt(size)
    run.font.bold = True
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT


def build_docx(report, tables: Dict[str, Dict[str, Any]], rules: ReportRules) -> bytes:
    """Собрать DOCX-отчёт по правилам."""
    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.left_margin = Cm(1.5)
    section.right_margin = Cm(1.5)

    _add_heading(document, rules.title, size=14)

    if rules.show_profile:
        for line in profile_lines(report, rules):
            paragraph = document.add_paragraph()
            run = paragraph.add_run(line)
            run.font.name = CELL_FONT
            run.font.size = Pt(10)

    for table in tables.values():
        columns: List[Column] = table["table"].columns
        rows: List[List[str]] = table["rows"]
        header_rows = _header_rows(table)
        document.add_paragraph()
        _add_heading(document, table["title"], size=12)
        word_table = document.add_table(rows=len(header_rows) + len(rows), cols=len(columns))
        table_rows = word_table.rows
        for row_offset, header_row in enumerate(header_rows):
            header_cells = table_rows[row_offset].cells
            for header_cell in header_row:
                col = header_cell.col
                target = header_cells[col]
                target.text = header_cell.text
                if header_cell.colspan > 1 or header_cell.rowspan > 1:
                    target.merge(table_rows[row_offset + header_cell.rowspan - 1].cells[col + header_cell.colspan - 1])
        for row_index, row in enumerate(rows, start=len(header_rows)):
            data_cells = table_rows[row_index].cells
            for column_index, value in enumerate(row):
                data_cells[column_index].text = value
        _style_table(word_table, header_height=len(header_rows))

    output = BytesIO()
    document.save(output)
    return output.getvalue()
