#!/usr/bin/env python3
"""Импорт источников рабочего каталога в канонические справочники и seed-данные.

Источники (только чтение, ничего не изменяется):

* ``справочники/бдму.xlsx`` — эталон корреляции: (УБИ, объект) -> СП, последствия,
  тактики, техники;
* ``справочники/ИС и негативные последствия.xlsx`` — виды риска У1/У2/У3,
  последствия уX.Y и их применимость по типам ИС;
* ``справочники/объекты.txt``, ``виды воздействия.txt``, ``способы реализации.txt``,
  ``нарушители.txt``, ``информационные системы.txt`` — словари;
* ``справочники/тактики_техники.csv`` — фактически XLSX: тактики Т1-Т10 и техники;
* ``сводная таблица.docx`` — уровни возможностей нарушителей по каждой УБИ;
* ``УБИ-объект.docx`` — наименования УБИ и источник угрозы;
* ``data/lib/Виды нарушителей.docx`` — категории (внешний/внутренний) и цели нарушителей;
* ``data/lib/Уровни возможностей нарушителей.docx`` — описания уровней Н1-Н4;
* ``data/lib/Виды рисков.docx`` — наименования видов риска У1-У3;
* ``data/lib/Способы реализации угроз безопасности информации.docx`` — связь
  «доступные интерфейсы -> способы реализации».

Результат:

* ``data/reference/*.yaml`` — канонические справочники;
* ``data/seed/threats_full.json`` — канонический корпус УБИ.

Запуск::

    python -m scripts.import_sources            # из каталога threat-modeler
    python -m scripts.import_sources --check     # только проверка, без записи
"""

from __future__ import annotations

import argparse
import io
import json
import re
import sys
from collections import Counter, OrderedDict, defaultdict
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple

try:
    import yaml
except ImportError:  # pragma: no cover
    print("Требуется PyYAML: pip install pyyaml")
    raise

try:
    import docx
    import openpyxl
except ImportError:  # pragma: no cover
    print("Требуются python-docx и openpyxl")
    raise


# --------------------------------------------------------------------------- #
# Расположение источников
# --------------------------------------------------------------------------- #

PROJECT_DIR = Path(__file__).resolve().parent.parent          # threat-modeler/
REPO_DIR = PROJECT_DIR.parent                                  # корень репозитория

SOURCES = {
    "bdmu": REPO_DIR / "справочники" / "бдму.xlsx",
    "is_consequences": REPO_DIR / "справочники" / "ИС и негативные последствия.xlsx",
    "objects_txt": REPO_DIR / "справочники" / "объекты.txt",
    "impacts_txt": REPO_DIR / "справочники" / "виды воздействия.txt",
    "methods_txt": REPO_DIR / "справочники" / "способы реализации.txt",
    "consequences_txt": REPO_DIR / "справочники" / "негативные последствия.txt",
    "consequence_table_docx": PROJECT_DIR / "data" / "lib" / "Негативные последствия.docx",
    "violators_txt": REPO_DIR / "справочники" / "нарушители.txt",
    "systems_txt": REPO_DIR / "справочники" / "информационные системы.txt",
    "tactics_xlsx": REPO_DIR / "справочники" / "тактики_техники.csv",
    "levels_xlsx": REPO_DIR / "Уровни возможностей нарушителей по УБИ.xlsx",
    "summary_docx": REPO_DIR / "сводная таблица.docx",
    "ubi_object_docx": REPO_DIR / "УБИ-объект.docx",
    "violator_types_docx": PROJECT_DIR / "data" / "lib" / "Виды нарушителей.docx",
    "violator_levels_docx": PROJECT_DIR / "data" / "lib" / "Уровни возможностей нарушителей.docx",
    "risk_types_docx": PROJECT_DIR / "data" / "lib" / "Виды рисков.docx",
    "interfaces_docx": PROJECT_DIR / "data" / "lib" / "Способы реализации угроз безопасности информации.docx",
}

REFERENCE_DIR = PROJECT_DIR / "data" / "reference"
SEED_DIR = PROJECT_DIR / "data" / "seed"


# --------------------------------------------------------------------------- #
# Соответствие «технология -> объекты» (выводится из состава объектов бдму)
# --------------------------------------------------------------------------- #

TECHNOLOGY_OBJECTS: "OrderedDict[str, Dict[str, Any]]" = OrderedDict([
    ("cloud", {"name": "Облачные технологии", "objects": ["О19", "О20"]}),
    ("virtualization", {"name": "Виртуализация", "objects": ["О5", "О6", "О9"]}),
    ("docker", {"name": "Docker / контейнеры", "objects": ["О1", "О21"]}),
    ("grid", {"name": "Грид-системы", "objects": ["О10", "О38"]}),
    ("mobile", {"name": "Мобильные устройства", "objects": ["О16"]}),
    ("ics", {"name": "Промышленные системы (ICS/SCADA)", "objects": ["О24"]}),
    ("smartcard", {"name": "Смарт-карты и аппаратная аутентификация", "objects": ["О32"]}),
    ("ml", {"name": "Искусственный интеллект и машинное обучение", "objects": ["О17", "О37"]}),
    ("bigdata", {"name": "Большие данные", "objects": ["О36"]}),
    ("supercomputer", {"name": "Суперкомпьютеры", "objects": ["О8", "О35"]}),
    ("removable_media", {"name": "Съёмные носители информации", "objects": ["О13", "О18"]}),
])
# Беспроводные каналы отдельным технологическим признаком не задаются: в корпусе
# нет отдельного объекта, признак выражается выбором интерфейса на шаге 3.

# Правила вывода «доступный интерфейс -> способы реализации» по документу
# «Способы реализации угроз безопасности информации.docx» (столбцы
# «Доступные интерфейсы» и «Способы реализации»).
INTERFACE_RULES: List[Tuple[str, str, str, Tuple[str, ...]]] = [
    (
        "external_network",
        "Внешние сетевые интерфейсы",
        "Интерфейсы взаимодействия с сетью «Интернет» и смежными системами/сетями",
        ("внешние сетевые интерфейсы",),
    ),
    (
        "internal_network",
        "Внутренние сетевые интерфейсы",
        "Интерфейсы взаимодействия с компонентами систем и сетей изнутри",
        ("внутренние сетевые интерфейсы",),
    ),
    (
        "user",
        "Интерфейсы для пользователей",
        "Пользовательские интерфейсы, в том числе веб-интерфейсы",
        ("интерфейсы для пользователей", "пользовательские веб-интерфейсы"),
    ),
    (
        "remote_access",
        "Интерфейсы удалённого доступа",
        "RDP, SSH, VNC, VPN и иные интерфейсы удалённого доступа",
        ("интерфейсы удалённого доступа", "интерфейсы удаленного доступа"),
    ),
    (
        "wireless",
        "Беспроводные интерфейсы",
        "Wi-Fi, Bluetooth, NFC и иные беспроводные каналы",
        ("беспроводные",),
    ),
    (
        "web",
        "Веб-интерфейсы",
        "HTTP/HTTPS-интерфейсы приложений и порталов",
        ("веб-интерфейсы",),
    ),
    (
        "removable_media",
        "Интерфейсы съёмных носителей и периферии",
        "Порты и интерфейсы для съёмных машинных носителей и периферийного оборудования",
        ("съемных машинных носителей", "съёмных машинных носителей"),
    ),
    (
        "maintenance",
        "Интерфейсы установки, настройки, испытаний",
        "Интерфейсы администрирования, управления, обслуживания, пусконаладочных работ",
        ("установки, настройки, испытаний",),
    ),
    (
        "supply_chain",
        "Доступ к поставляемым и обслуживаемым компонентам",
        "Доступ к компонентам при поставке, обслуживании и ремонте в сторонних организациях",
        ("поставляемым или находящимся на обслуживании",),
    ),
    (
        "physical",
        "Физический доступ к оборудованию",
        "Непосредственный доступ к оборудованию при обслуживании",
        ("физического доступа к оборудованию",),
    ),
]

VIOLATOR_TYPE_CATEGORY_FALLBACK = {
    "специальные службы иностранных государств": "external",
    "террористические, экстремистские группировки": "external",
    "преступные группы (криминальные структуры)": "external",
    "отдельные физические лица (хакеры)": "external",
    "бывшие (уволенные) работники (пользователи)": "external",
    "конкурирующие организации": "external",
    "разработчики программных, программно-аппаратных средств": "internal",
    "лица, обеспечивающие поставку программных, программно-аппаратных средств, обеспечивающих систем": "external",
    "лица, привлекаемые для установки, настройки, испытаний, пусконаладочных и иных видов работ": "internal",
    "поставщики услуг связи": "internal",
    "лица, обеспечивающие функционирование систем и сетей или обеспечивающих систем оператора (администрация, охрана, уборщики и др.)": "internal",
    "авторизованные пользователи систем и сетей": "internal",
    "системные администраторы и администраторы безопасности": "internal",
}

# Расхождения формулировок между «нарушители.txt» и «Виды нарушителей.docx»
VIOLATOR_TYPE_ALIASES = {
    "поставщики услуг связи": "поставщики вычислительных услуг, услуг связи",
    "лица, обеспечивающие поставку программных, программно-аппаратных средств, обеспечивающих систем": (
        "лица, обеспечивающие поставку программных, программно-аппаратных средств, обеспечивающих систем"
    ),
    "лица, обеспечивающие функционирование систем и сетей или обеспечивающих систем оператора "
    "(администрация, охрана, уборщики и др.)": (
        "лица, обеспечивающие функционирование систем и сетей или обеспечивающие систему оператора "
        "(администрация, охрана, уборщики и т.д.)"
    ),
}

# Расхождения формулировок между «нарушители.txt» и
# «Уровни возможностей нарушителей по УБИ.xlsx»
VIOLATOR_LEVEL_ALIASES = {
    "преступные группы (криминальные структуры)": (
        "преступные группы (два лица и более, действующие по единому плану)"
    ),
    "отдельные физические лица (хакеры)": "физическое лицо (хакер)",
    "бывшие (уволенные) работники (пользователи)": "бывшие работники (пользователи)",
    "поставщики услуг связи": "поставщики вычислительных услуг, услуг связи",
    "лица, обеспечивающие функционирование систем и сетей или обеспечивающих систем оператора "
    "(администрация, охрана, уборщики и др.)": (
        "лица, обеспечивающие функционирование систем и сетей или обеспечивающих систем "
        "(администрация, охрана, уборщики и т.д.)"
    ),
}


def normalize_name(value: str) -> str:
    """Нормализовать название для сопоставления: регистр, пробелы, конечная пунктуация."""
    text = normalize_ws(value or "").strip().rstrip(";.").strip()
    return text.lower()


# --------------------------------------------------------------------------- #
# Утилиты
# --------------------------------------------------------------------------- #

def read_text(path: Path) -> str:
    raw = path.read_bytes()
    for enc in ("utf-8", "cp1251"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def normalize_ws(value: str) -> str:
    return re.sub(r"\s+", " ", value or "").strip()


def norm_code(prefix: str, value: str) -> Optional[str]:
    """Нормализовать код вида «СП1», «Т1.4», «п2.19», «УБИ.7»."""
    m = re.search(prefix + r"\s*(\d+)(?:\.(\d+))?", value, re.I)
    if not m:
        return None
    if m.group(2) is None:
        return f"{prefix}{int(m.group(1))}"
    return f"{prefix}{int(m.group(1))}.{int(m.group(2))}"


def split_codes(value: Any, prefix: str) -> List[str]:
    """Разобрать строку «СП1, СП2, СП8» в список нормализованных кодов."""
    text = "" if value is None else str(value)
    out: List[str] = []
    pattern = re.compile(prefix + r"\s*(\d+)(?:\.(\d+))?", re.I)
    for m in pattern.finditer(text):
        code = f"{prefix}{int(m.group(1))}" + (f".{int(m.group(2))}" if m.group(2) else "")
        if code not in out:
            out.append(code)
    return out


def threat_id(value: str) -> Optional[str]:
    m = re.search(r"УБИ\s*\.?\s*(\d+)", value or "", re.I)
    return f"УБИ.{int(m.group(1)):03d}" if m else None


def normalize_consequence_code(value: str) -> Optional[str]:
    """«п2.19» и «у2.19» — одна и та же позиция; канон — «у2.19»."""
    m = re.search(r"[пу]\s*(\d+)\s*\.\s*(\d+)", value or "", re.I)
    return f"у{int(m.group(1))}.{int(m.group(2))}" if m else None


def slugify(value: str) -> str:
    translit = {
        "а": "a", "б": "b", "в": "v", "г": "g", "д": "d", "е": "e", "ё": "e", "ж": "zh",
        "з": "z", "и": "i", "й": "y", "к": "k", "л": "l", "м": "m", "н": "n", "о": "o",
        "п": "p", "р": "r", "с": "s", "т": "t", "у": "u", "ф": "f", "х": "h", "ц": "c",
        "ч": "ch", "ш": "sh", "щ": "sch", "ъ": "", "ы": "y", "ь": "", "э": "e", "ю": "yu",
        "я": "ya",
    }
    out = []
    for ch in normalize_ws(value).lower():
        if ch in translit:
            out.append(translit[ch])
        elif ch.isalnum():
            out.append(ch)
        elif ch in " -—/(),.":
            out.append("_")
    slug = re.sub(r"_+", "_", "".join(out)).strip("_")
    return slug[:48] or "item"


def docx_rows(path: Path):
    document = docx.Document(str(path))
    for table in document.tables:
        for row in table.rows:
            yield [normalize_ws(cell.text) for cell in row.cells]


# --------------------------------------------------------------------------- #
# Парсеры словарей
# --------------------------------------------------------------------------- #

class Issues:
    def __init__(self) -> None:
        self.items: List[str] = []

    def add(self, message: str) -> None:
        if message not in self.items:
            self.items.append(message)

    def report(self) -> None:
        if not self.items:
            print("  замечаний нет")
            return
        print(f"  замечаний: {len(self.items)}")
        for item in self.items:
            print("    !", item)


def parse_impacts(issues: Issues) -> Dict[str, Dict[str, str]]:
    text = read_text(SOURCES["impacts_txt"])
    impacts: "OrderedDict[str, Dict[str, str]]" = OrderedDict()
    for line in text.splitlines():
        m = re.match(r"^(В\d+)\s+(.*)$", normalize_ws(line))
        if not m:
            continue
        code = m.group(1)
        impacts[code] = {"code": code, "name": normalize_ws(m.group(2)).rstrip(";.")}
    if len(impacts) != 7:
        issues.add(f"виды воздействия: ожидалось 7, получено {len(impacts)}")
    return impacts


def parse_methods(issues: Issues) -> Dict[str, Dict[str, str]]:
    text = read_text(SOURCES["methods_txt"])
    methods: "OrderedDict[str, Dict[str, str]]" = OrderedDict()
    for line in text.splitlines():
        m = re.match(r"^(СП\d+)\s+(.*)$", normalize_ws(line))
        if not m:
            continue
        code = m.group(1)
        methods[code] = {"code": code, "name": normalize_ws(m.group(2)).rstrip(";.")}
    if len(methods) != 9:
        issues.add(f"способы реализации: ожидалось 9, получено {len(methods)}")
    return methods


def parse_system_types(issues: Issues) -> Dict[str, Dict[str, str]]:
    names = [normalize_ws(x) for x in read_text(SOURCES["systems_txt"]).splitlines() if normalize_ws(x)]
    descriptions: Dict[str, str] = {}
    workbook = openpyxl.load_workbook(SOURCES["is_consequences"], read_only=True, data_only=True)
    sheet = workbook["Входные данные"]
    for row in sheet.iter_rows(values_only=True):
        if not row or not row[0]:
            continue
        key = normalize_ws(str(row[0]))
        if key:
            descriptions[key] = normalize_ws(str(row[1] or ""))
    types: "OrderedDict[str, Dict[str, str]]" = OrderedDict()
    for name in names:
        types[name] = {"code": name, "name": name, "description": descriptions.get(name, "")}
    if len(types) != 14:
        issues.add(f"типы ИС: ожидалось 14, получено {len(types)}")
    missing_desc = [k for k, v in types.items() if not v["description"]]
    if missing_desc:
        issues.add(f"типы ИС без описания: {missing_desc}")
    return types


def parse_consequences(issues: Issues) -> Tuple[Dict[str, Dict[str, Any]], Dict[str, Dict[str, str]]]:
    """Собрать канонический справочник последствий.

    Нумерация берётся из ``бдму.xlsx``: файл ссылается одновременно на ``п2.6`` и
    ``п2.26``, значит исходный перечень содержит 26 позиций для У2. В
    ``негативные последствия.txt`` метка ``п2.6`` «слиплас» внутрь текста ``п2.5``,
    а в ``ИС и негативные последствия.xlsx`` эту позицию объединили и перенумеровали
    остаток (``у2.6`` файла = ``п2.7`` исходного перечня). Поэтому:

    * ``п1.N`` / ``п3.N`` -> ``у1.N`` / ``у3.N`` без сдвига;
    * ``п2.1``..``п2.5`` -> ``у2.1``..``у2.5``;
    * ``п2.6`` -> восстанавливается из текста ``п2.5`` (применимость берётся у ``у2.5``);
    * ``п2.7``..``п2.26`` -> ``у2.6``..``у2.25``.
    """
    risk_names: Dict[str, str] = {}
    for row in docx_rows(SOURCES["risk_types_docx"]):
        if row and re.match(r"^У[123]$", row[0]) and len(row) >= 2:
            risk_names[row[0]] = row[1]
    risk_names.setdefault("У1", "Ущерб физическому лицу")
    risk_names.setdefault("У2", "Ущерб юридическому лицу, индивидуальному предпринимателю")
    risk_names.setdefault("У3", "Ущерб государству")

    canonical = parse_consequences_txt(issues)
    xlsx = parse_consequences_xlsx(issues)

    consequences: "OrderedDict[str, Dict[str, Any]]" = OrderedDict()
    for code, item in canonical.items():
        risk, number = code.split(".")[0], int(code.split(".")[1])
        if risk == "у2" and number >= 7:
            source_code = f"у2.{number - 1}"
            source_note = "сдвиг нумерации в ИС и негативные последствия.xlsx"
        elif risk == "у2" and number == 5:
            source_code = "у2.5"
            source_note = "позиция объединена в источнике с у2.6"
        elif risk == "у2" and number == 6:
            source_code = "у2.5"
            source_note = "позиция объединена в источнике; применимость принята по у2.5"
        else:
            source_code = code
            source_note = ""
        source = xlsx.get(source_code)
        if source is None:
            issues.add(f"последствие {code}: нет строки {source_code} в ИС и негативные последствия.xlsx")
            continue
        if not source_note and normalize_ws(source["name"]).lower() != normalize_ws(item["name"]).lower():
            issues.add(
                f"последствие {code}: текст не совпадает с источником "
                f"({item['name'][:40]!r} != {source['name'][:40]!r})"
            )
        consequences[code] = {
            "code": code,
            "risk": source["risk"] or {"у1": "У1", "у2": "У2", "у3": "У3"}[risk],
            "name": item["name"],
            "conditions": source["conditions"],
            "system_types": source["system_types"],
            "source_code": source_code,
            "source_note": source_note,
        }

    extra = set(xlsx) - {c["source_code"] for c in consequences.values()}
    if extra:
        issues.add(f"неиспользованные строки справочника последствий: {sorted(extra)}")
    risks = {code: {"code": code, "name": name} for code, name in sorted(risk_names.items())}
    return consequences, risks


def parse_consequences_txt(issues: Issues) -> "OrderedDict[str, Dict[str, Any]]":
    """Разобрать ``негативные последствия.txt`` и восстановить потерянную метку п2.6."""
    text = read_text(SOURCES["consequences_txt"])
    items: "OrderedDict[str, Dict[str, Any]]" = OrderedDict()
    for raw_line in text.splitlines():
        line = normalize_ws(raw_line)
        if not line:
            continue
        m = re.match(r"^([пу]\s*\d+\.\d+)\s*(.*)$", line, re.I)
        if not m:
            issues.add(f"последствия: строка без кода {line[:60]!r}")
            continue
        code = normalize_consequence_code(m.group(1))
        body = m.group(2)
        inner = re.search(r"[пу]\s*(\d+\.\d+)", body, re.I)
        if code and inner:
            # Внутри текста «слиплас» следующая позиция — разделяем её.
            first_code = code
            second_code = normalize_consequence_code(inner.group(0))
            first_name = body[: inner.start()].strip().rstrip("(,").strip()
            first_name = re.sub(r"\(в том числе закупка$", "", first_name).strip()
            tail = body[inner.end():].strip().lstrip(")").strip()
            if tail and tail[0].islower():
                tail = "Необходимость дополнительных (незапланированных) затрат на закупку " + tail
            items[first_code] = {"code": first_code, "name": first_name}
            if second_code:
                items[second_code] = {"code": second_code, "name": normalize_ws(tail).rstrip(".")}
            continue
        if code:
            items[code] = {"code": code, "name": body.rstrip(".")}
    if not items:
        issues.add("последствия: не удалось разобрать негативные последствия.txt")
    return items


def parse_consequences_xlsx(issues: Issues) -> Dict[str, Dict[str, Any]]:
    """Прочитать применимость последствий по типам ИС."""
    workbook = openpyxl.load_workbook(SOURCES["is_consequences"], read_only=True, data_only=True)
    sheet = workbook["Лист1"]
    rows = list(sheet.iter_rows(values_only=True))
    result: Dict[str, Dict[str, Any]] = {}
    for row in rows[1:]:
        if not row or not row[1]:
            continue
        code = normalize_consequence_code(str(row[1]))
        if not code:
            issues.add(f"последствие с нераспознанным кодом: {row[1]!r}")
            continue
        conditions = normalize_ws(str(row[3] or ""))
        result[code] = {
            "code": code,
            "risk": normalize_ws(str(row[0] or "")),
            "name": normalize_ws(str(row[2] or "")),
            "conditions": conditions,
            "system_types": parse_conditions(conditions),
        }
    return result


def parse_conditions(value: str) -> List[str]:
    """Выделить типы ИС из строки «Условия принадлежности»."""
    known = [
        "ГИС избирательных комиссий",
        "ИС обороны и безопасности",
        "портал госуслуг (ЕПГУ, РПГУ)",
        "ИС внешнеэкономической деятельности",
        "ГИС МЧС",
        "ГИС МИД",
        "КОРП ИС",
        "веб-портал",
        "веб-ресурс",
        "АСУ ТП",
        "ИСПДн",
        "ПДн",
        "ГИС",
        "КИИ",
        "КТ",
        "ДСП",
    ]
    found: List[str] = []
    lowered = value.lower()
    for name in known:
        if name.lower() in lowered and name not in found:
            found.append(name)
    # «веб-ресурс» в источнике — синоним веб-портала
    if "веб-ресурс" in found and "веб-портал" not in found:
        found[found.index("веб-ресурс")] = "веб-портал"
    if "ИСПДн" in found and "ПДн" not in found:
        found.append("ПДн")
    return found


def parse_objects(issues: Issues) -> "OrderedDict[str, Dict[str, Any]]":
    text = read_text(SOURCES["objects_txt"])
    objects: "OrderedDict[str, Dict[str, Any]]" = OrderedDict()
    for line in text.splitlines():
        m = re.match(r"^(\d+)\s+(.*)$", normalize_ws(line))
        if not m:
            continue
        num = int(m.group(1))
        code = f"О{num}"
        objects[code] = {"code": code, "num": num, "name": normalize_ws(m.group(2))}
    if len(objects) != 38:
        issues.add(f"объекты: ожидалось 38, получено {len(objects)}")
    return objects


def parse_violator_types(levels: Dict[str, Dict[str, Any]], issues: Issues) -> "OrderedDict[str, Dict[str, Any]]":
    raw_names = [normalize_ws(x) for x in read_text(SOURCES["violators_txt"]).splitlines() if normalize_ws(x)]
    details: Dict[str, Dict[str, str]] = {}
    for row in docx_rows(SOURCES["violator_types_docx"]):
        if len(row) < 3 or not row[1]:
            continue
        if row[2] not in ("Внешний", "Внутренний"):
            continue
        details[normalize_name(row[1])] = {
            "category": "external" if row[2] == "Внешний" else "internal",
            "goals": row[3] if len(row) > 3 else "",
        }

    # «тип нарушителя -> уровень возможностей» по файлу уровней
    level_by_violator: Dict[str, str] = {}
    for level_code, item in levels.items():
        for name in item["violator_types"]:
            level_by_violator[normalize_name(name)] = level_code

    types: "OrderedDict[str, Dict[str, Any]]" = OrderedDict()
    for raw in raw_names:
        display = normalize_ws(raw).rstrip(";.").strip()
        display = display[:1].upper() + display[1:]
        key = normalize_name(raw)
        lookup = VIOLATOR_TYPE_ALIASES.get(key, key)
        info = details.get(lookup, {})
        category = info.get("category") or VIOLATOR_TYPE_CATEGORY_FALLBACK.get(key, "external")
        code = slugify(raw)
        while code in types:
            code += "_2"
        level_key = VIOLATOR_LEVEL_ALIASES.get(key, key)
        level = level_by_violator.get(level_key, "")
        if not level:
            issues.add(f"вид нарушителя «{display}»: не найден уровень возможностей в файле уровней")
        types[code] = {
            "code": code,
            "name": display,
            "category": category,
            "level": level,
            "goals": info.get("goals", ""),
        }
    if len(types) != 13:
        issues.add(f"виды нарушителей: ожидалось 13, получено {len(types)}")
    unmatched = [normalize_name(name) for name in raw_names
                 if VIOLATOR_TYPE_ALIASES.get(normalize_name(name), normalize_name(name)) not in details]
    if unmatched:
        issues.add(f"виды нарушителей без категории в Виды нарушителей.docx: {unmatched}")
    internal = [code for code, item in types.items() if item["category"] == "internal"]
    external = [code for code, item in types.items() if item["category"] == "external"]
    if not internal:
        issues.add("виды нарушителей: не определено ни одного внутреннего нарушителя")
    by_level = Counter(item["level"] for item in types.values())
    print(
        f"    нарушители: внешних {len(external)}, внутренних {len(internal)}; "
        f"по уровням {dict(sorted(by_level.items()))}"
    )
    return types


def parse_violator_levels(issues: Issues) -> "OrderedDict[str, Dict[str, Any]]":
    levels: "OrderedDict[str, Dict[str, Any]]" = OrderedDict()
    for row in docx_rows(SOURCES["violator_levels_docx"]):
        if not row or not re.match(r"^Н[1-4]$", row[0]):
            continue
        code = row[0]
        if code in levels:
            continue
        levels[code] = {
            "code": code,
            "rank": int(code[1]),
            "name": row[1] if len(row) > 1 else "",
            "description": row[2] if len(row) > 2 else "",
            "violators": row[3] if len(row) > 3 else "",
        }
    if len(levels) != 4:
        issues.add(f"уровни нарушителей: ожидалось 4, получено {len(levels)}")
    return levels


def merge_levels(
    levels: "OrderedDict[str, Dict[str, Any]]",
    capability: Dict[str, Dict[str, Any]],
    violator_types: Dict[str, Dict[str, Any]],
    issues: Issues,
) -> "OrderedDict[str, Dict[str, Any]]":
    """Дополнить уровни данными файла «Уровни возможностей нарушителей по УБИ.xlsx»."""
    code_by_level: Dict[str, List[str]] = {}
    for code, item in violator_types.items():
        code_by_level.setdefault(item.get("level") or "", []).append(code)

    for code, item in levels.items():
        extra = capability.get(code)
        if not extra:
            issues.add(f"уровень {code}: нет данных в файле уровней возможностей")
            continue
        item["potential"] = extra["potential"]
        item["methods_extended"] = extra["methods_extended"]
        item["level_violators"] = code_by_level.get(code, [])
        if extra["name"] and not item.get("name"):
            item["name"] = extra["name"]

        # категории из файла уровней не должны противоречить «Виды нарушителей.docx»
        for name, category in extra["categories"].items():
            key = VIOLATOR_LEVEL_ALIASES.get(name, name)
            for type_code in code_by_level.get(code, []):
                if normalize_name(violator_types[type_code]["name"]) == key:
                    if violator_types[type_code]["category"] != category:
                        issues.add(
                            f"вид нарушителя «{violator_types[type_code]['name']}»: категория "
                            f"в файле уровней ({category}) не совпадает с «Виды нарушителей.docx» "
                            f"({violator_types[type_code]['category']})"
                        )
    return levels


def parse_capability_levels(issues: Issues) -> "OrderedDict[str, Dict[str, Any]]":
    """Разобрать «Уровни возможностей нарушителей по УБИ.xlsx».

    Файл даёт для каждого уровня возможностей Н1-Н4: потенциал, перечень видов
    нарушителей и перечень способов реализации, доступных нарушителю этого уровня.
    ВНИМАНИЕ: в файле используется расширенная нумерация способов реализации
    (СП.1 ... СП.26), которая не совпадает с нумерацией СП1-СП9 методики и
    ``бдму.xlsx``, поэтому перечень сохраняется отдельно (``methods_extended``)
    и не смешивается с кодами корреляции.
    """
    if not SOURCES["levels_xlsx"].exists():
        issues.add(f"источник не найден: {SOURCES['levels_xlsx']}")
        return OrderedDict()

    workbook = openpyxl.load_workbook(SOURCES["levels_xlsx"], data_only=True)
    sheet = workbook.worksheets[0]
    rows = list(sheet.iter_rows(values_only=True))
    header = [normalize_ws(str(c or "")) for c in rows[0]]
    expected = ["Потенциал", "№", "Уровень возможностей", "Наименование", "Тип нарушителя", "Способ реализации"]
    if header[:6] != expected:
        issues.add(f"уровни возможностей: неожидаемые заголовки {header[:6]}")

    levels: "OrderedDict[str, Dict[str, Any]]" = OrderedDict()
    current: Optional[Dict[str, Any]] = None
    for row in rows[1:]:
        values = list(row) + [""] * 6
        # Ячейку со способами реализации не нормализуем: в ней несколько строк
        potential = normalize_ws("" if values[0] is None else str(values[0]))
        number = normalize_ws("" if values[1] is None else str(values[1]))
        level_name = normalize_ws("" if values[2] is None else str(values[2]))
        violator = normalize_ws("" if values[3] is None else str(values[3]))
        category = normalize_ws("" if values[4] is None else str(values[4]))
        methods_raw = "" if values[5] is None else str(values[5])
        if number:
            if not re.match(r"^Н[1-4]$", number):
                issues.add(f"уровни возможностей: неизвестный уровень {number!r}")
                continue
            current = levels.setdefault(
                number,
                {
                    "code": number,
                    "rank": int(number[1]),
                    "potential": potential,
                    "name": level_name,
                    "methods_extended": [],
                    "violator_types": [],
                    "categories": {},
                },
            )
            if potential:
                current["potential"] = potential
            if level_name:
                current["name"] = level_name
            # В источнике столбец «Потенциал» заполнен не для всех уровней
            if not current.get("potential"):
                lowered = (level_name or "").lower()
                for keyword, value in (
                    ("базовыми повышенными", "Средний повышенный"),
                    ("средними", "Средний"),
                    ("высокими", "Высокий"),
                    ("базовыми", "Низкий"),
                ):
                    if keyword in lowered:
                        current["potential"] = value
                        current["potential_inferred"] = True
                        break
            for line in str(methods_raw).splitlines():
                line = normalize_ws(line)
                match = re.match(r"^(СП\.?\d+)\s+(.*)$", line)
                if match:
                    current["methods_extended"].append(
                        {"code": match.group(1), "name": normalize_ws(match.group(2))}
                    )
        if current is None:
            continue
        if violator:
            if violator not in current["violator_types"]:
                current["violator_types"].append(violator)
            if category:
                current["categories"][normalize_name(violator)] = (
                    "external" if category.lower().startswith("внеш") else "internal"
                )

    if len(levels) != 4:
        issues.add(f"уровни возможностей: ожидалось 4, получено {len(levels)}")
    total_violators = sum(len(item["violator_types"]) for item in levels.values())
    if total_violators != 13:
        issues.add(f"уровни возможностей: ожидалось 13 видов нарушителей, получено {total_violators}")
    for item in levels.values():
        if not item["methods_extended"]:
            issues.add(f"уровень {item['code']}: не указаны способы реализации")
    return levels


def parse_tactics(issues: Issues) -> "OrderedDict[str, Dict[str, Any]]":
    data = SOURCES["tactics_xlsx"].read_bytes()
    workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheet = workbook.worksheets[0]
    rows = list(sheet.iter_rows(values_only=True))
    tactics: "OrderedDict[str, Dict[str, Any]]" = OrderedDict()
    for row in rows[1:]:
        if not row or not row[0]:
            continue
        tactic_code = norm_code("Т", str(row[0]))
        technique_code = norm_code("Т", str(row[2] or ""))
        if not tactic_code or not technique_code:
            issues.add(f"тактики/техники: нераспознанная строка {row[:3]!r}")
            continue
        tactic = tactics.setdefault(
            tactic_code, {"code": tactic_code, "name": normalize_ws(str(row[1] or "")), "techniques": OrderedDict()}
        )
        tactic["techniques"][technique_code] = {
            "code": technique_code,
            "name": normalize_ws(str(row[3] or "")).split(":")[0][:120],
            "description": normalize_ws(str(row[3] or "")),
        }
    if len(tactics) != 10:
        issues.add(f"тактики: ожидалось 10, получено {len(tactics)}")
    return tactics


def parse_interfaces(methods: Dict[str, Dict[str, str]], issues: Issues) -> Dict[str, Dict[str, Any]]:
    source_rows: List[Tuple[str, str, str]] = []
    for row in docx_rows(SOURCES["interfaces_docx"]):
        if len(row) < 6:
            continue
        available, sposoby = row[4], row[5]
        codes = split_codes(sposoby, "СП")
        if available and codes:
            source_rows.append((available, sposoby, ",".join(codes)))

    interfaces: "OrderedDict[str, Dict[str, Any]]" = OrderedDict()
    for code, name, description, keywords in INTERFACE_RULES:
        found: List[str] = []
        for available, _sposoby, codes in source_rows:
            low = available.lower()
            if any(keyword in low for keyword in keywords):
                for c in codes.split(","):
                    if c and c not in found:
                        found.append(c)
        unknown = [c for c in found if c not in methods]
        if unknown:
            issues.add(f"интерфейс {code}: неизвестные СП {unknown}")
        interfaces[code] = {
            "code": code,
            "name": name,
            "description": description,
            "methods": sorted(found, key=lambda x: int(x[2:])),
            "source": (
                "data/lib/Способы реализации угроз безопасности информации.docx "
                "(столбцы «Доступные интерфейсы» и «Способы реализации»)"
            ),
        }
        if not found:
            issues.add(f"интерфейс {code}: не выведено ни одного способа реализации")
    return interfaces


# --------------------------------------------------------------------------- #
# Корпус УБИ
# --------------------------------------------------------------------------- #

def parse_threats(
    objects: Dict[str, Dict[str, Any]],
    methods: Dict[str, Dict[str, str]],
    consequences: Dict[str, Dict[str, Any]],
    tactics: Dict[str, Dict[str, Any]],
    issues: Issues,
) -> List[Dict[str, Any]]:
    workbook = openpyxl.load_workbook(SOURCES["bdmu"], read_only=True, data_only=True)
    sheet = workbook.worksheets[0]
    rows = [r for r in sheet.iter_rows(values_only=True)]
    header = [normalize_ws(str(c or "")) for c in rows[0]]
    expected = ["УБИ", "Объект", "Способы реализации", "Негативные последствия", "Тактики", "Основные техники"]
    if header[:6] != expected:
        issues.add(f"бдму: неожидаемые заголовки {header[:6]}")

    object_by_num = {obj["num"]: code for code, obj in objects.items()}
    techniques_by_code: Dict[str, str] = {}
    for tactic in tactics.values():
        for code, technique in tactic["techniques"].items():
            techniques_by_code[code] = tactic["code"]

    threats: "OrderedDict[str, Dict[str, Any]]" = OrderedDict()
    for row in rows[1:]:
        if not row or not row[0]:
            continue
        tid = threat_id(str(row[0]))
        if not tid:
            issues.add(f"бдму: нераспознанный идентификатор УБИ {row[0]!r}")
            continue
        object_field = normalize_ws(str(row[1] or ""))
        m = re.match(r"^(\d+)\s+(.*)$", object_field)
        if not m:
            issues.add(f"{tid}: нераспознанный объект {object_field!r}")
            continue
        obj_code = object_by_num.get(int(m.group(1)))
        if not obj_code:
            issues.add(f"{tid}: объект №{m.group(1)} отсутствует в справочнике")
            continue

        pair_methods = [c for c in split_codes(row[2], "СП") if c in methods]
        pair_consequences = []
        for token in re.findall(r"[пу]\s*\d+\s*\.\s*\d+", str(row[3] or ""), re.I):
            code = normalize_consequence_code(token)
            if code and code in consequences and code not in pair_consequences:
                pair_consequences.append(code)
        pair_tactics = [c for c in split_codes(row[4], "Т") if c in tactics and "." not in c]
        pair_techniques = [c for c in split_codes(row[5], "Т") if c in techniques_by_code]

        unknown_consequences = [
            t for t in re.findall(r"[пу]\s*\d+\s*\.\s*\d+", str(row[3] or ""), re.I)
            if normalize_consequence_code(t) not in consequences
        ]
        if unknown_consequences:
            issues.add(f"{tid}: неизвестные последствия {sorted(set(unknown_consequences))}")

        threat = threats.setdefault(
            tid,
            {"id": tid, "name": "", "source_note": "", "objects_hint": [], "pairs": [], "program": []},
        )
        if any(p["object"] == obj_code for p in threat["pairs"]):
            issues.add(f"{tid}: дублирующая пара с объектом {obj_code}")
        threat["pairs"].append(
            {
                "object": obj_code,
                "methods": pair_methods,
                "consequences": pair_consequences,
                "tactics": pair_tactics,
                "techniques": pair_techniques,
            }
        )

    # наименования и источник угрозы
    for row in docx_rows(SOURCES["ubi_object_docx"]):
        if len(row) < 4:
            continue
        tid = threat_id(row[0])
        if not tid:
            continue
        threat = threats.setdefault(
            tid, {"id": tid, "name": "", "source_note": "", "objects_hint": [], "pairs": [], "program": []}
        )
        threat["name"] = threat["name"] or row[1]
        threat["source_note"] = threat["source_note"] or row[2]
        threat["objects_hint"] = [normalize_ws(x) for x in re.split(r"[,;]", row[3]) if normalize_ws(x)]

    # уровни нарушителей
    levels_seen = set()
    for row in docx_rows(SOURCES["summary_docx"]):
        if not row or not row[0].startswith("УБИ"):
            continue
        tid = threat_id(row[0])
        if not tid:
            continue
        threat = threats.get(tid)
        if threat is None:
            issues.add(f"{tid}: есть в сводной таблице, но отсутствует в бдму")
            continue
        if tid in levels_seen:
            continue
        levels_seen.add(tid)
        m = re.match(r"^УБИ\.\d+\s*:\s*(.*)$", row[0])
        if m and m.group(1):
            threat["name"] = threat["name"] or normalize_ws(m.group(1))
        threat["violator_int"] = parse_level_list(row[1] if len(row) > 1 else "")
        threat["violator_ext"] = parse_level_list(row[2] if len(row) > 2 else "")
        threat["note"] = row[8] if len(row) > 8 else ""

    without_levels = [t["id"] for t in threats.values() if "violator_int" not in t]
    if without_levels:
        issues.add(f"без уровней нарушителей: {len(without_levels)} УБИ, например {without_levels[:5]}")
    without_names = [t["id"] for t in threats.values() if not t["name"]]
    if without_names:
        issues.add(f"без наименования: {len(without_names)} УБИ, например {without_names[:5]}")
    empty_pairs = [t["id"] for t in threats.values() if not t["pairs"]]
    if empty_pairs:
        issues.add(f"без связей с объектами: {len(empty_pairs)} УБИ")

    result = list(threats.values())
    result.sort(key=lambda t: int(t["id"].split(".")[1]))
    return result


def parse_level_list(value: str) -> List[str]:
    if not value or value.strip() in ("―", "-", "—"):
        return []
    return [f"Н{int(m)}" for m in re.findall(r"Н\s*(\d)", value)]


# --------------------------------------------------------------------------- #
# Соответствие «негативное последствие -> виды воздействия»
# --------------------------------------------------------------------------- #

def parse_consequence_table(issues: Issues) -> List[Dict[str, Any]]:
    """Прочитать таблицу 3 Методики: П1-П17 -> объекты и виды воздействия."""
    rows: List[Dict[str, Any]] = []
    for row in docx_rows(SOURCES["consequence_table_docx"]):
        if len(row) < 3:
            continue
        match = re.match(r"^(П\d+)\.\s*(.*)$", row[0])
        if not match:
            continue
        rows.append(
            {
                "code": match.group(1),
                "name": match.group(2),
                "objects": re.findall(r"О\d+", row[1]),
                "impacts": re.findall(r"В\d+", row[2]),
            }
        )
    if len(rows) != 17:
        issues.add(f"таблица 3: ожидалось 17 строк последствий, получено {len(rows)}")
    return rows


def name_tokens(value: str) -> str:
    """Нормализовать название последствия для сопоставления."""
    text = normalize_ws(value).lower()
    text = re.sub(r"^[пу]\d+\.\s*", "", text)
    text = re.sub(r"\(незапланированных\)", "", text)
    text = re.sub(r"[^0-9a-zа-яё ]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def align_consequence_table(
    table_rows: List[Dict[str, Any]], consequences: Dict[str, Dict[str, Any]], issues: Issues
) -> Dict[str, Dict[str, Any]]:
    """Сопоставить П1-П17 со списком последствий риск-группы У2.

    Таблица 3 Методики содержит укрупнённые последствия П1-П17, а рабочий корпус —
    расширенный перечень у2.1-у2.26. Оба перечня упорядочены одинаково, поэтому
    выполняется монотонное выравнивание (динамическое программирование) с
    максимизацией схожести названий; результат проверяется на порядок.
    """
    import difflib

    у2 = [code for code in consequences if code.startswith("у2.")]
    у2.sort(key=lambda code: int(code.split(".")[1]))
    if not у2:
        issues.add("не найдено последствий риск-группы У2 для сопоставления с таблицей 3")
        return {}

    rows = table_rows
    n, m = len(rows), len(у2)
    sim = [[0.0] * m for _ in range(n)]
    for i, row in enumerate(rows):
        left = name_tokens(row["name"])
        for j, code in enumerate(у2):
            right = name_tokens(consequences[code]["name"])
            sim[i][j] = difflib.SequenceMatcher(None, left, right).ratio()

    # dp[i][j] — лучшая сумма схожести для первых i строк таблицы на первых j позициях у2
    neg = float("-inf")
    dp = [[neg] * (m + 1) for _ in range(n + 1)]
    back: List[List[Optional[str]]] = [[None] * (m + 1) for _ in range(n + 1)]
    for j in range(m + 1):
        dp[0][j] = 0.0
    for i in range(1, n + 1):
        for j in range(1, m + 1):
            best, choice = dp[i][j - 1], "skip"
            if dp[i - 1][j - 1] > neg:
                candidate = dp[i - 1][j - 1] + sim[i - 1][j - 1]
                if candidate > best:
                    best, choice = candidate, "match"
            dp[i][j] = best
            back[i][j] = choice

    mapping: Dict[str, Dict[str, Any]] = {}
    i, j = n, m
    pairs: List[Tuple[int, int]] = []
    while i > 0 and j > 0:
        if back[i][j] == "match":
            pairs.append((i - 1, j - 1))
            i, j = i - 1, j - 1
        else:
            j -= 1
    pairs.reverse()

    for row_index, col_index in pairs:
        code = у2[col_index]
        score = sim[row_index][col_index]
        mapping[code] = {
            "code": code,
            "methodology_code": rows[row_index]["code"],
            "methodology_name": rows[row_index]["name"],
            "impacts": rows[row_index]["impacts"],
            "objects": rows[row_index]["objects"],
            "match": round(score, 2),
        }
        if score < 0.6:
            issues.add(
                f"сопоставление таблицы 3: {rows[row_index]['code']} -> {code} "
                f"с низкой схожестью {score:.2f}"
            )

    unmatched = [row["code"] for index, row in enumerate(rows) if index not in {p[0] for p in pairs}]
    if unmatched:
        issues.add(f"таблица 3: не сопоставлены строки {unmatched}")
    impacts_all = {impact for row in rows for impact in row["impacts"]}
    print(
        f"    таблица 3: сопоставлено {len(mapping)} последствий, "
        f"виды воздействия: {', '.join(sorted(impacts_all, key=lambda x: int(x[1:])))}"
    )
    return mapping


def build_consequence_impacts(
    consequences: Dict[str, Dict[str, Any]],
    mapping: Dict[str, Dict[str, Any]],
) -> Tuple[Dict[str, List[str]], List[str]]:
    """Вернуть «последствие -> виды воздействия» и перечень несопоставленных позиций."""
    impacts: Dict[str, List[str]] = {}
    for code, item in mapping.items():
        if code in consequences and item["impacts"]:
            impacts[code] = sorted(item["impacts"], key=lambda x: int(x[1:]))
    unmatched = [code for code in consequences if code not in impacts]
    return impacts, unmatched


# --------------------------------------------------------------------------- #
# Запись результатов
# --------------------------------------------------------------------------- #

def plain(value: Any) -> Any:
    """Преобразовать OrderedDict в обычные структуры для YAML/JSON."""
    if isinstance(value, dict):
        return {k: plain(v) for k, v in value.items()}
    if isinstance(value, (list, tuple)):
        return [plain(v) for v in value]
    return value


def dump_yaml(path: Path, payload: Dict[str, Any], header: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", encoding="utf-8") as handle:
        handle.write(f"# {header}\n")
        handle.write("# Файл сгенерирован scripts/import_sources.py — не редактировать вручную.\n\n")
        yaml.safe_dump(plain(payload), handle, allow_unicode=True, sort_keys=False, width=120)


def build(check_only: bool) -> int:
    issues = Issues()
    print("Импорт источников")
    for name, path in SOURCES.items():
        if not path.exists():
            issues.add(f"источник не найден: {name} -> {path}")

    impacts = parse_impacts(issues)
    methods = parse_methods(issues)
    system_types = parse_system_types(issues)
    consequences, risk_types = parse_consequences(issues)
    objects = parse_objects(issues)
    capability_levels = parse_capability_levels(issues)
    violator_types = parse_violator_types(capability_levels, issues)
    violator_levels = parse_violator_levels(issues)
    violator_levels = merge_levels(violator_levels, capability_levels, violator_types, issues)
    tactics = parse_tactics(issues)
    interfaces = parse_interfaces(methods, issues)
    threats = parse_threats(objects, methods, consequences, tactics, issues)

    consequence_table = parse_consequence_table(issues)
    alignment = align_consequence_table(consequence_table, consequences, issues)
    consequence_impacts, consequences_without_impacts = build_consequence_impacts(consequences, alignment)

    used_objects = {p["object"] for t in threats for p in t["pairs"]}
    unused_objects = sorted(set(objects) - used_objects, key=lambda c: objects[c]["num"])
    unused_consequences = sorted(set(consequences) - {c for t in threats for p in t["pairs"] for c in p["consequences"]})

    used_by_corpus = {c for t in threats for p in t["pairs"] for c in p["consequences"]}
    threats_with_impacts = [
        t for t in threats
        if any(c in consequence_impacts for p in t["pairs"] for c in p["consequences"])
    ]

    print(f"  УБИ: {len(threats)}, пар (УБИ, объект): {sum(len(t['pairs']) for t in threats)}")
    print(f"  объекты: {len(objects)} (задействовано {len(used_objects)}, не задействовано {len(unused_objects)})")
    print(f"  последствия: {len(consequences)} (задействовано {len(consequences) - len(unused_consequences)})")
    print(f"  способы реализации: {len(methods)}, виды воздействия: {len(impacts)}")
    print(
        f"  виды воздействия выведены для {len(consequence_impacts)} последствий; "
        f"УБИ с видами воздействия: {len(threats_with_impacts)} из {len(threats)}"
    )
    print(f"  тактики: {len(tactics)}, техники: {sum(len(t['techniques']) for t in tactics.values())}")
    print(f"  виды нарушителей: {len(violator_types)}, уровни: {len(violator_levels)}")
    extended_total = sum(len(item.get("methods_extended") or []) for item in violator_levels.values())
    print(
        f"    расширенный перечень способов реализации из файла уровней: {extended_total} позиций "
        f"(нумерация СП.N не совпадает с СП1-СП9 методики — хранится отдельно)"
    )
    print(f"  интерфейсы: {len(interfaces)}")
    for code, item in interfaces.items():
        print(f"    {code}: {len(item['methods'])} СП -> {item['methods']}")
    print("Замечания по источникам:")
    issues.report()

    if check_only:
        return 1 if issues.items else 0

    reference: Dict[str, Tuple[Dict[str, Any], str]] = {
        "objects.yaml": ({"objects": objects}, "Справочник объектов воздействия (38 позиций, бдму.xlsx)"),
        "impacts.yaml": ({"impacts": impacts}, "Справочник видов воздействия (В1-В7, Методика ФСТЭК)"),
        "methods.yaml": ({"methods": methods}, "Справочник способов реализации (СП1-СП9, Методика ФСТЭК)"),
        "consequences.yaml": (
            {"consequences": consequences, "risk_types": risk_types},
            "Справочник негативных последствий (уX.Y) и видов риска (У1-У3)",
        ),
        "consequence_impacts.yaml": (
            {
                "consequence_impacts": consequence_impacts,
                "mapping": alignment,
                "without_impacts": consequences_without_impacts,
                "used_by_corpus": sorted(used_by_corpus, key=lambda c: (int(c[1]), int(c.split(".")[1]))),
            },
            "Соответствие «негативное последствие -> виды воздействия» (таблица 3 Методики)",
        ),
        "system_types.yaml": ({"system_types": system_types}, "Типы информационных систем"),
        "violators.yaml": (
            {"violator_types": violator_types, "violator_levels": violator_levels},
            "Виды нарушителей и уровни возможностей Н1-Н4",
        ),
        "tactics.yaml": ({"tactics": tactics}, "Тактики (Т1-Т10) и техники"),
        "interfaces.yaml": (
            {"interfaces": interfaces},
            "Типы интерфейсов и выведенные из Методики способы реализации",
        ),
        "technologies.yaml": (
            {"technologies": TECHNOLOGY_OBJECTS},
            "Технологические признаки системы и соответствующие объекты воздействия",
        ),
    }
    for filename, (payload, header) in reference.items():
        dump_yaml(REFERENCE_DIR / filename, payload, header)
        print(f"  записан data/reference/{filename}")

    seed = {
        "generated_by": "scripts/import_sources.py",
        "objects": {code: obj["name"] for code, obj in objects.items()},
        "threats": [
            {
                "id": t["id"],
                "name": t["name"],
                "source_note": t.get("source_note", ""),
                "note": t.get("note", ""),
                "violator_int": t.get("violator_int", []),
                "violator_ext": t.get("violator_ext", []),
                "pairs": t["pairs"],
            }
            for t in threats
        ],
    }
    SEED_DIR.mkdir(parents=True, exist_ok=True)
    with open(SEED_DIR / "threats_full.json", "w", encoding="utf-8") as handle:
        json.dump(plain(seed), handle, ensure_ascii=False, indent=1)
    print(f"  записан data/seed/threats_full.json ({len(threats)} УБИ)")
    return 1 if issues.items else 0


def main(argv: Optional[Sequence[str]] = None) -> int:
    parser = argparse.ArgumentParser(description="Импорт источников рабочего каталога")
    parser.add_argument("--check", action="store_true", help="только проверка источников, без записи")
    args = parser.parse_args(argv)
    return build(args.check)


if __name__ == "__main__":
    sys.exit(main())
