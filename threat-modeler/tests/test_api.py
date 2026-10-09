"""Интеграционные тесты API: wizard, предпросмотр, экспорт XLSX/DOCX."""

from __future__ import annotations

import io

from conftest import extract_csrf

XLSX_MAGIC = b"PK\x03\x04"
DOCX_MAGIC = b"PK\x03\x04"


def fill_wizard(client, system_types=None, technologies=None, objects=None, interfaces=None, impacts=None):
    """Пройти шаги 1-4 и вернуть ответ последнего шага."""
    from app.core.reference_loader import ReferenceLoader

    system_types = system_types or list(ReferenceLoader.get_system_types())[:2]
    technologies = technologies or []
    objects = objects or ["О4", "О31", "О23"]
    violator_types = ReferenceLoader.get_violator_types()
    external = next(code for code, item in violator_types.items() if item["category"] == "external")
    internal = next(code for code, item in violator_types.items() if item["category"] == "internal")

    page = client.get("/api/wizard/")
    assert page.status_code == 200
    csrf = extract_csrf(page.text)

    response = client.post(
        "/api/wizard/step1",
        data={
            "csrf_token": csrf,
            "system_types": system_types,
            "technologies": technologies,
            "processes_pd": "on",
            "pd_security_level": "2",
        },
        follow_redirects=False,
    )
    assert response.status_code == 303, response.text

    page = client.get("/api/wizard/step/2")
    csrf = extract_csrf(page.text)
    response = client.post(
        "/api/wizard/step2",
        data={"csrf_token": csrf, "external_types": [external], "internal_types": [internal]},
        follow_redirects=False,
    )
    assert response.status_code == 303, response.text

    page = client.get("/api/wizard/step/3")
    csrf = extract_csrf(page.text)
    response = client.post(
        "/api/wizard/step3",
        data={"csrf_token": csrf},
        follow_redirects=False,
    )
    assert response.status_code == 303, response.text

    page = client.get("/api/wizard/step/4")
    csrf = extract_csrf(page.text)
    return client.post(
        "/api/wizard/step4",
        data={
            "csrf_token": csrf,
            "selected_objects": objects,
        },
        follow_redirects=False,
    )


# --------------------------------------------------------------------------- #
# Базовые маршруты
# --------------------------------------------------------------------------- #

def test_health(client):
    """Health-check отвечает."""
    response = client.get("/health")
    assert response.status_code == 200
    assert response.json()["status"] == "healthy"


def test_root_redirects_to_wizard(client):
    """Корень перенаправляет на wizard."""
    response = client.get("/", follow_redirects=False)
    assert response.status_code == 303
    assert response.headers["location"].startswith("/api/wizard")


def test_wizard_steps_render(client):
    """Все пять шагов отдаются без ошибок."""
    for step in range(1, 6):
        response = client.get(f"/api/wizard/step/{step}")
        assert response.status_code == 200, f"шаг {step}: {response.status_code}"


def test_unknown_step_returns_404(client):
    """Несуществующий шаг — 404."""
    assert client.get("/api/wizard/step/9").status_code == 404


def test_reference_endpoints(client):
    """Справочники отдаются и содержат ожидаемое количество позиций."""
    payload = client.get("/api/reference/").json()
    counts = payload["counts"]
    assert counts["objects"] == 38
    assert counts["methods"] == 9
    assert counts["system_types"] == 13
    assert counts["interfaces"] == 10
    assert counts["consequences"] == 67
    assert counts["violator_types"] == 13
    assert counts["violator_levels"] == 4
    assert counts["tactics"] == 10

    assert len(client.get("/api/reference/system-types").json()) == 13

    violators = client.get("/api/reference/violators").json()
    assert len(violators["violator_types"]) == 13
    for level in violators["violator_levels"].values():
        assert level["methods"], f"уровень {level['code']} без способов реализации"
    levels = {item["level"] for item in violators["violator_types"].values()}
    assert levels == {"Н1", "Н2", "Н3", "Н4"}
    assert violators["violator_goals_table"], "нет таблицы целей нарушителей"


# --------------------------------------------------------------------------- #
# Опрос
# --------------------------------------------------------------------------- #

def test_step1_validation_error_keeps_form(client):
    """Ошибка валидации возвращает форму с сообщением, а не JSON-ошибку."""
    page = client.get("/api/wizard/")
    csrf = extract_csrf(page.text)
    response = client.post("/api/wizard/step1", data={"csrf_token": csrf})
    assert response.status_code == 400
    assert "Проверьте введённые данные" in response.text


def test_csrf_is_enforced(client):
    """POST без корректного CSRF-токена отклоняется."""
    client.get("/api/wizard/")
    response = client.post("/api/wizard/step1", data={"system_types": ["ПДн"]})
    assert response.status_code == 403


def test_full_wizard_flow_and_report_page(client):
    """Сквозной проход wizard: страница отчёта содержит обе таблицы и данные."""
    response = fill_wizard(client)
    assert response.status_code == 303, response.text

    page = client.get("/api/wizard/step/5")
    assert page.status_code == 200
    assert "Таблица 1" in page.text
    assert "Таблица 2" in page.text
    assert "УБИ." in page.text, "в отчёте нет строк с УБИ"
    assert "Виды воздействия" in page.text


def test_report_page_without_profile_is_empty_but_renders(client):
    """Без пройденного опроса страница отчёта открывается и не падает."""
    page = client.get("/api/wizard/step/5")
    assert page.status_code == 200
    assert "Таблица 1" in page.text


def test_clear_wipes_session_and_shows_clean_form(client):
    """Кнопка «Очистить данные» полностью сбрасывает опрос."""
    fill_wizard(client)
    page = client.get("/api/wizard/step/5")
    csrf = extract_csrf(page.text)
    cleared = client.post("/api/wizard/clear", data={"csrf_token": csrf}, follow_redirects=False)
    assert cleared.status_code == 303

    fresh = client.get("/api/wizard/")
    assert fresh.status_code == 200
    # ни один тип ИС и ни один объект не отмечен
    assert "checked>" not in fresh.text
    # и данные шага 2 больше не подставляются
    step2 = client.get("/api/wizard/step/2")
    assert 'value="Н3" selected' not in step2.text


def test_step5_page_contains_tables(client):
    """Страница отчёта содержит таблицы 1 и 2."""
    fill_wizard(client)
    page = client.get("/api/wizard/step/5")
    assert page.status_code == 200
    assert "Таблица 1" in page.text
    assert "Таблица 2" in page.text
    assert "Перечень возможных" in page.text


# --------------------------------------------------------------------------- #
# Экспорт
# --------------------------------------------------------------------------- #

def test_export_xlsx(client):
    """XLSX-отчёт скачивается и является корректным архивом Office."""
    fill_wizard(client)
    response = client.get("/api/reports/export.xlsx")
    assert response.status_code == 200
    assert response.content[:4] == XLSX_MAGIC
    assert "attachment" in response.headers["content-disposition"]

    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(response.content))
    assert workbook.sheetnames == [f"Таблица {i}" for i in range(1, 8)]
    sheet = workbook["Таблица 7"]
    assert sheet.max_row > 3, "таблица 7 не содержит строк с УБИ"
    # шапка соответствует правилам отчёта
    assert sheet.cell(row=2, column=1).value == "Идентификатор УБИ"
    assert sheet.cell(row=2, column=2).value == "Уровень возможностей нарушителей"


def test_export_docx(client):
    """DOCX-отчёт скачивается и открывается python-docx."""
    fill_wizard(client)
    response = client.get("/api/reports/export.docx")
    assert response.status_code == 200
    assert response.content[:4] == DOCX_MAGIC

    import docx

    document = docx.Document(io.BytesIO(response.content))
    assert len(document.tables) == 7
    assert any("Таблица 1" in paragraph.text for paragraph in document.paragraphs)


def test_export_alias_without_extension(client):
    """Маршрут /api/reports/export отдаёт XLSX."""
    fill_wizard(client)
    response = client.get("/api/reports/export")
    assert response.status_code == 200
    assert response.content[:4] == XLSX_MAGIC


def test_export_respects_report_rules(client, tmp_path):
    """Изменение report-rules.txt меняет и предпросмотр, и выгрузку."""
    from app.config import settings

    custom = tmp_path / "report-rules.txt"
    custom.write_text(
        "[report]\ntitle = Свой отчёт\nshow_profile = no\n"
        "[table1]\nenabled = no\n\n"
        "[table7]\ntitle = Только идентификаторы\ncolumn = ubi_id | Код УБИ | 16\n",
        encoding="utf-8",
    )
    old_data_dir = settings.data_dir
    try:
        settings.data_dir = str(tmp_path)
        fill_wizard(client)
        page = client.get("/api/wizard/step/5")
        assert page.status_code == 200
        assert "Свой отчёт" in page.text
        assert "Таблица 1" not in page.text
        assert "Код УБИ" in page.text

        export = client.get("/api/reports/export.xlsx")
        assert export.status_code == 200
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(export.content))
        assert workbook.sheetnames == ["Таблица 7"]
        assert workbook["Таблица 7"].cell(row=2, column=1).value == "Код УБИ"
    finally:
        settings.data_dir = old_data_dir


def test_broken_rules_file_is_reported_on_page(client, tmp_path):
    """Ошибка в report-rules.txt показывается на странице, а не ломает сервис."""
    from app.config import settings

    (tmp_path / "report-rules.txt").write_text(
        "[table2]\ncolumn = нет_такого_поля | X\n", encoding="utf-8"
    )
    old_data_dir = settings.data_dir
    try:
        settings.data_dir = str(tmp_path)
        fill_wizard(client)
        page = client.get("/api/wizard/step/5")
        assert page.status_code == 200
        assert "Правила отчёта не применены" in page.text

        export = client.get("/api/reports/export.xlsx")
        assert export.status_code == 500
        assert "неизвестное поле" in export.json()["detail"]
    finally:
        settings.data_dir = old_data_dir


def test_export_without_profile_is_rejected(client):
    """Без профиля экспорт недоступен."""
    client.get("/api/wizard/")
    assert client.get("/api/reports/export.xlsx").status_code == 400
